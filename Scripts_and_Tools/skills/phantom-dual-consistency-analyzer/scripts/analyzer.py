import os
import sys
import json
import re
import argparse
import base64
from datetime import datetime
from docx import Document
import xlsxwriter

def parse_docx(path):
    if not os.path.exists(path): return None
    try:
        doc = Document(path)
        res = {'filename': os.path.basename(path), 'heading_map': {}, 'full_text': ''}
        txts = []
        cur_h = "START"
        res['heading_map'][cur_h] = []
        for p in doc.paragraphs:
            t = p.text.strip()
            if not t: continue
            txts.append(t)
            is_h = 'Heading' in (p.style.name if p.style else '') or bool(re.match(r'^[壹貳參肆伍陸柒捌玖拾一二三四五六七八九十\d（(]+[、.)\s]', t))
            if is_h and len(t) < 60:
                cur_h = t
                res['heading_map'][cur_h] = []
            else:
                res['heading_map'][cur_h].append(t)
        res['full_text'] = '\n'.join(txts)
        return res
    except: return None

def search_sec(doc, name, kws):
    if not doc: return ""
    for h, c in doc['heading_map'].items():
        if name in h: return "\n".join(c)
    for kw in kws:
        if kw in doc['full_text']:
            idx = doc['full_text'].index(kw)
            return doc['full_text'][max(0, idx-100):idx+800]
    return ""

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--plan", required=True)
    parser.add_argument("--rfp", required=True)
    parser.add_argument("--output", default="consistency_report.xlsx")
    args = parser.parse_args()

    plan_d = parse_docx(args.plan)
    rfp_d = parse_docx(args.rfp)

    # Simplified logic for the skill script
    dims = [
        {"dim": "採購定位", "kw": ["定位", "目標", "背景"]},
        {"dim": "數據轉譯", "kw": ["轉譯", "白話", "圖文", "回饋"]},
        {"dim": "技術規格", "kw": ["規格", "功能", "RPM", "軌跡"]},
        {"dim": "MIT要求", "kw": ["產地", "製造", "MIT", "臺灣"]},
        {"dim": "資安規範", "kw": ["資安", "保護", "個資", "弱點"]}
    ]

    wb = xlsxwriter.Workbook(args.output)
    ws = wb.add_worksheet('分析報告')
    header_f = wb.add_format({'bold':True, 'bg_color':'#DDEAF6', 'border':1})
    cell_f = wb.add_format({'border':1, 'text_wrap':True, 'valign':'top'})

    ws.write_row(0, 0, ["檢核維度", "計畫要點(節錄)", "招標落實(節錄)", "比對結論", "建議修正"], header_f)
    
    for i, d in enumerate(dims):
        p_txt = search_sec(plan_d, d['dim'], d['kw'])
        r_txt = search_sec(rfp_d, d['dim'], d['kw'])
        ws.write(i+1, 0, d['dim'], cell_f)
        ws.write(i+1, 1, p_txt[:1000], cell_f)
        ws.write(i+1, 2, r_txt[:1000], cell_f)
        ws.write(i+1, 3, "自動比對：一致" if (p_txt[:10] == r_txt[:10] and p_txt) else "需人工覆核", cell_f)
        ws.write(i+1, 4, "請確認是否符合核定計畫需求", cell_f)

    wb.close()
    print(f"Report saved: {args.output}")

if __name__ == "__main__": main()
